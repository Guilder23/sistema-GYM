from django.shortcuts import render
from django.contrib.auth.decorators import login_required
from django.http import HttpResponse
from django.db.models import Sum, Count
from django.db.models.functions import TruncMonth, TruncDate
from apps.clients.models import Client
from apps.memberships.models import Payment, Membership, MembershipPlan
from apps.access.models import AccessRecord
from apps.inventory.models import Sale, Product
from datetime import datetime, timedelta
import csv
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill


@login_required
def reports_dashboard(request):
    today = datetime.now().date()
    month_ago = today - timedelta(days=30)
    
    # Basic Stats
    total_revenue_month = Payment.objects.filter(date__gte=month_ago).aggregate(Sum('amount'))['amount__sum'] or 0
    active_clients = Client.objects.filter(is_active=True).count()
    inactive_clients = Client.objects.filter(is_active=False).count()
    today_access = AccessRecord.objects.filter(timestamp__date=today).count()
    
    # Revenue Trend (Last 6 Months)
    six_months_ago = today - timedelta(days=180)
    revenue_trend = Payment.objects.filter(date__gte=six_months_ago) \
        .annotate(month=TruncMonth('date')) \
        .values('month') \
        .annotate(total=Sum('amount')) \
        .order_by('month')
    
    # Membership Distribution
    plan_distribution = Membership.objects.filter(active=True) \
        .values('plan__name') \
        .annotate(count=Count('id')) \
        .order_by('-count')
    
    # Attendance Trend (Last 7 Days)
    seven_days_ago = today - timedelta(days=7)
    attendance_trend = AccessRecord.objects.filter(timestamp__date__gte=seven_days_ago) \
        .annotate(day=TruncDate('timestamp')) \
        .values('day') \
        .annotate(count=Count('id')) \
        .order_by('day')

    # Prepare data for Chart.js
    revenue_labels = [r['month'].strftime('%b %Y') for r in revenue_trend]
    revenue_data = [float(r['total']) for r in revenue_trend]
    
    plan_labels = [p['plan__name'] for p in plan_distribution]
    plan_data = [p['count'] for p in plan_distribution]
    
    attendance_labels = [a['day'].strftime('%d/%m') for a in attendance_trend]
    attendance_data = [a['count'] for a in attendance_trend]
    
    context = {
        'total_revenue_month': total_revenue_month,
        'active_clients': active_clients,
        'inactive_clients': inactive_clients,
        'today_access': today_access,
        'revenue_labels': revenue_labels,
        'revenue_data': revenue_data,
        'plan_labels': plan_labels,
        'plan_data': plan_data,
        'attendance_labels': attendance_labels,
        'attendance_data': attendance_data,
    }
    return render(request, 'reports/dashboard.html', context)


@login_required
def export_clients_excel(request):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Clientes"
    
    # Styles
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="4F46E5", end_color="4F46E5", fill_type="solid")
    center_aligned = Alignment(horizontal="center")
    
    headers = ['ID', 'Nombre', 'Apellido', 'Documento', 'Email', 'Teléfono', 'Estado', 'Fecha Registro']
    ws.append(headers)
    
    for cell in ws[1]:
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = center_aligned
        
    clients = Client.objects.all().order_by('last_name')
    for client in clients:
        ws.append([
            client.id,
            client.first_name,
            client.last_name,
            client.ci,
            client.email,
            client.phone,
            'Activo' if client.is_active else 'Inactivo',
            client.join_date.strftime('%d/%m/%Y')
        ])
        
    response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    response['Content-Disposition'] = 'attachment; filename=clientes_gym.xlsx'
    wb.save(response)
    return response


from django.template.loader import get_template
from xhtml2pdf import pisa


@login_required
def export_payments_pdf(request):
    payments = Payment.objects.all().order_by('-date')
    template_path = 'reports/pdf_payments.html'
    
    total = sum(p.amount for p in payments) if payments.exists() else 0
    
    context = {
        'payments': payments,
        'today': datetime.now(),
        'total': total
    }
    
    response = HttpResponse(content_type='application/pdf')
    response['Content-Disposition'] = 'attachment; filename="reporte_pagos.pdf"'
    
    template = get_template(template_path)
    html = template.render(context)
    
    pisa_status = pisa.CreatePDF(html, dest=response)
    if pisa_status.err:
       return HttpResponse('Error al generar el PDF', status=500)
    return response


@login_required
def export_access_csv(request):
    response = HttpResponse(content_type='text/csv')
    response['Content-Disposition'] = 'attachment; filename="historial_accesos.csv"'
    
    writer = csv.writer(response)
    writer.writerow(['Fecha y Hora', 'Cliente', 'Documento', 'Válido', 'Nota'])
    
    records = AccessRecord.objects.all().order_by('-timestamp')
    for r in records:
        writer.writerow([
            r.timestamp.strftime('%d/%m/%Y %H:%M:%S'),
            r.client.full_name,
            r.client.ci,
            'Sí' if r.valid else 'No',
            r.message
        ])
        
    return response
